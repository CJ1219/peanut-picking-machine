from pathlib import Path
import shutil
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from build_hardware_report import ROOT, OUT, para, heading, picture, table, font_run

SOURCE = Path(r'C:\Users\user\OneDrive - 中原大學\桌面\大學作品集\花生機\決賽繳交資料夾\02_照片')
OUT = ROOT / '硬體與配電成果報告.docx'
ASSETS = ROOT / '硬體報告_qa' / 'assets'
ASSETS.mkdir(exist_ok=True)
for name in ['花生機2.PNG', '花生機.PNG', 'LINE_ALBUM_機器照片_260907_1.jpg', 'LINE_ALBUM_機器照片_260907_2.jpg', 'LINE_ALBUM_機器照片_260907_5.jpg']:
    target = ASSETS / name
    if not target.exists():
        shutil.copy2(SOURCE / name, target)

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Mm(12.7)
s.footer_distance = Mm(6)
n = doc.styles['Normal']
n.font.name = 'Times New Roman'
n.font.size = Pt(12)
n.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '標楷體')
n.paragraph_format.line_spacing = Pt(20)
f = s.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.CENTER
font_run(f.add_run('硬體與配電成果報告　'), 9)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
f._p.append(fld)

def page(title):
    doc.add_page_break()
    heading(doc, title)

para(doc, '2026 中興大學專題實作競賽', 13, align=WD_ALIGN_PARAGRAPH.CENTER)
p = para(doc, '花生智慧辨識與自動分選系統', 20, True, WD_ALIGN_PARAGRAPH.CENTER)
p.style = doc.styles['Title']
para(doc, '硬體與配電成果報告', 16, True, WD_ALIGN_PARAGRAPH.CENTER)
picture(doc, ASSETS / 'LINE_ALBUM_機器照片_260907_1.jpg', '圖 1　整機實體與下方配電配置', Mm(98))
para(doc, '本作品以生產線末端的花生品質辨識與分選為應用情境，自行設計滾筒輸送機構、撥片分流結構、機架及控制箱，將影像辨識與實體分選動作整合於同一設備。透過摩擦驅動使滾筒自轉，配合四通道伺服撥片與集中式配電配置，呈現由機構設計、試作改良到機電整合的實作成果。')

page('一 整機配置與設計範圍')
para(doc, '整機由上方的滾筒輸送模組、末端撥片與漏斗、直立控制桿及下方配電空間組成。相機安裝於辨識區上方；控制桿集中配置操作按鈕及方向選擇開關，方便操作員控制設備。底部腳輪可供移動與鎖定。')
picture(doc, ASSETS / '花生機2.PNG', '圖 2　整機 3D 設計圖與分流結構配置', Mm(135))
table(doc, [['項目', '設計規格'], ['滾筒模組尺寸', '約長 40 × 寬 35 × 高 20 cm'], ['整機尺寸', '含控制桿與機箱約長 60 × 寬 40 × 高 120 cm'], ['輸送驅動', '17HS4417 步進馬達 2 顆，同軸配置'], ['分流驅動', 'MG996R 伺服馬達 4 顆'], ['相機位置', '辨識區正上方，以環境光拍攝'], ['辨識至分流距離', '辨識區中心至撥片約 20 cm']])
para(doc, '花生投入後，利用相鄰滾筒形成的凹槽引導成列排列，並隨輸送機構前進至辨識區。滾筒同時提供承載、輸送與翻轉作用，讓花生在移動過程中呈現不同表面，供上方相機擷取影像。')

page('二 滾筒傳動改良歷程')
heading(doc, '初期齒輪與齒條方案', 2)
para(doc, '初期以環狀排列的滾筒搭配滾筒間小齒輪，使各滾筒彼此連動；並在辨識區下方設置齒條，當滾筒端部齒輪通過齒條時帶動滾筒自轉。設計目的在於讓花生翻轉，使相機能取得不同表面的影像。')
heading(doc, '傳動結構簡化與改良', 2)
para(doc, '滾筒以壓克力管製作，實際來料的內徑與外徑均較原設計尺寸大約 0.3 mm。此尺寸偏差除了影響滾筒組件的配合與齒輪嚙合，也改變整圈滾筒排列與間距。原設計規劃配置 23 根滾筒，實際裝配後調整為 22 根，整圈仍留有部分剩餘空間。')
para(doc, '為控制滾筒間距，我們採用八字形連接件連接相鄰滾筒，以連接件約束彼此的位置關係，管理調整根數後留下的空間。同時逐步縮短齒條並移除滾筒間小齒輪，將材料尺寸、排列間距與傳動結構一併納入改良。八字形連接件負責間距控制，辨識區下方的摩擦材料則負責帶動滾筒自轉，各自對應不同的機構需求。')
para(doc, '根據實際裝配及齒輪進入齒條時的接觸情形，進一步將滾筒自轉方式調整為摩擦驅動。此調整減少自轉機構對精確對齒的依賴，呈現團隊依材料尺寸與試作結果修正設計的能力。')
heading(doc, '目前摩擦驅動方案', 2)
para(doc, '最終取消齒條，在辨識區的滾筒下方固定羽毛球拍握把布。滾筒通過時，藉由與握把布接觸產生的摩擦力帶動自轉，減少齒輪嚙合環節。兩顆 17HS4417 步進馬達採同軸配置，設計目的為增加輸送驅動能力。')
para(doc, '此設計將輸送與滾筒自轉的作用分開處理：步進馬達負責輸送驅動，辨識區下方的摩擦材料負責帶動滾筒自轉。透過常見握把布材料實現接觸摩擦機構，也便於日後替換與維護。')
heading(doc, '分選作動與影像配置', 2)
para(doc, '辨識畫面以輔助線劃分四個通道，分別對應末端的四組撥片。每組撥片由一顆 MG996R 伺服馬達驅動，依辨識結果將不良花生導向撥片後方、機器下方的收集盒；良品則滑向前方漏斗，形成兩條收集路徑。')
para(doc, '辨識區中心距離撥片約 20 cm，讓影像擷取區與分流作動區具有明確的位置關係。滾筒排列、上方相機與末端撥片共同構成連續的辨識分選流程。')

page('三 控制箱與配線配置')
para(doc, '我們參考工業設備的配電配置方式，將控制元件與驅動元件集中於機架下方，以端子台作為接線集中點，搭配導軌與配線槽整理線路，便於查線、維修與更換元件。整機 3D 圖呈現配置規劃，實際裝配以實體照片為準。')
picture(doc, ASSETS / '花生機.PNG', '圖 3　3D 設計中的配電空間與元件配置', Mm(125))
para(doc, '實體控制箱採分層配置，上方為控制板與接線端子，中段配置端子台及繼電器，下方配置驅動器與金屬外殼電源供應器。線路透過端子台集中連接，搭配橫向與直向配線槽安排走線路徑，使不同功能的元件具有清楚的安裝位置。')
table(doc, [['配置', '設計用途'], ['端子台', '集中電源及控制線路接點，方便對照與維修'], ['導軌', '提供元件固定位置'], ['配線槽', '引導線材路徑並整理配線'], ['可鎖定腳輪', '移動設備及定位使用']])

page('四 電源與急停控制')
picture(doc, ASSETS / 'LINE_ALBUM_機器照片_260907_5.jpg', '圖 4　控制箱內部實體配線', Mm(105))
para(doc, '設備由 110 V 插座供電，經 100 W 電源供應器轉換為 12 V 輸出。電源控制配置 12 V／10 A 繼電器及急停開關；按下急停時，透過繼電器斷開電源，實現整機斷電控制。急停開關設於控制桿上方，便於操作員直接觸及。')
para(doc, '四顆 MG996R 伺服馬達分別驅動各通道的撥片。伺服供電由 12 V 電源經降壓模組轉換為 6 V，再供應 MG996R 使用。')

page('五 操作介面與維護')
picture(doc, ASSETS / 'LINE_ALBUM_機器照片_260907_2.jpg', '圖 5　控制桿與操作按鈕實體配置', Mm(85))
table(doc, [['操作元件', '功能與狀態'], ['綠色按鈕與指示', '啟動；運行時綠燈長亮'], ['紅色按鈕與指示', '停止；停止時紅燈長亮'], ['黃色按鈕與指示', '異常時黃燈閃爍；按下後清除可解除異常或重試連線'], ['方向選擇開關', '選擇輸送方向'], ['急停開關', '透過繼電器切斷整機電源']])
para(doc, '黃色按鈕清除異常時，若異常條件仍存在，例如生產設定尚未填妥，則應維持異常狀態。相機連線失敗時可藉此重新嘗試連線。')
para(doc, '維護時應檢查握把布的磨耗、髒污及固定狀況，並確認滾筒、撥片與收集位置沒有異物或偏移。這些項目用於持續確認摩擦接觸與分流路徑。')

heading(doc, '六 自製成果與整合特色')
para(doc, '滾筒機構、分流結構、機架、控制桿及機箱配置均由團隊自行設計；馬達、電子元件、驅動板與操作開關等採購現成零組件。作品的硬體成果包含機構試作、傳動問題排除與配電整合。')
para(doc, '本作品的機構特色在於利用滾筒間隙引導花生排列，並以接觸摩擦方式帶動自轉；分流端以四組伺服撥片對應辨識通道，將影像判定轉換為實際物料分流。配電端則整合端子台、導軌、配線槽、降壓供電與急停控制，將機構與控制元件集中於可移動的機架中。')

for t in doc.tables:
    t.autofit = False
    t.columns[0].width = Mm(43)
    t.columns[1].width = Mm(141)
    for idx, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for i, c in enumerate(row.cells):
            c.width = Mm(43 if i == 0 else 141)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = Pt(18)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e = OxmlElement('w:'+edge)
                for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    e.set(qn('w:'+k),v)
                borders.append(e)
            c._tc.get_or_add_tcPr().append(borders)
for inline in list(doc.element.body.iter(qn('wp:inline'))):
    anchor = OxmlElement('wp:anchor')
    for key, value in {'distT':'0', 'distB':'91440', 'distL':'0', 'distR':'0',
                       'simplePos':'0', 'relativeHeight':'0', 'behindDoc':'0',
                       'locked':'0', 'layoutInCell':'1', 'allowOverlap':'0'}.items():
        anchor.set(key, value)
    pos = OxmlElement('wp:simplePos')
    pos.set('x', '0')
    pos.set('y', '0')
    anchor.append(pos)
    horizontal = OxmlElement('wp:positionH')
    horizontal.set('relativeFrom', 'column')
    align = OxmlElement('wp:align')
    align.text = 'center'
    horizontal.append(align)
    anchor.append(horizontal)
    vertical = OxmlElement('wp:positionV')
    vertical.set('relativeFrom', 'paragraph')
    offset = OxmlElement('wp:posOffset')
    offset.text = '0'
    vertical.append(offset)
    anchor.append(vertical)
    for tag in ['wp:extent', 'wp:effectExtent']:
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement('wp:wrapTopAndBottom'))
    for child in list(inline):
        anchor.append(child)
    parent = inline.getparent()
    parent.replace(inline, anchor)
    paragraph = parent
    while paragraph is not None and paragraph.tag != qn('w:p'):
        paragraph = paragraph.getparent()
    if paragraph is not None:
        from docx.text.paragraph import Paragraph
        pf = Paragraph(paragraph, doc).paragraph_format
        pf.line_spacing = 1.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
doc.save(OUT)
shutil.copy2(OUT, ROOT / '決賽繳交資料夾' / OUT.name)
print('Saved report with', len(list(doc.element.body.iter(qn('wp:wrapTopAndBottom')))), 'top-and-bottom figures')
